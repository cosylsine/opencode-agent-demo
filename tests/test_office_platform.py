"""Real private storage, tools and HTTP boundaries; model reasoning is explicitly injected."""

import io
import time
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select

from apps.orchestrator.config import Settings
from apps.orchestrator.main import create_app
from packages.adapters.example import ExampleObservationSource
from packages.adapters.platform_store import users
from tests.fakes import FakeClient

ORIGIN = "http://127.0.0.1:18000"
PASSWORD = "test-password-2026"


def headers(client, key=None):
    import uuid

    auth = client.get("/api/v1/auth/me").json()
    return {
        "Origin": ORIGIN,
        "X-CSRF-Token": auth["csrf_token"],
        "Idempotency-Key": key or str(uuid.uuid4()),
    }


def login(client, username="admin", password=PASSWORD):
    response = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
        headers=headers(client),
    )
    assert response.status_code == 200, response.text
    return response.json()["user"]


def register(client, username):
    response = client.post(
        "/api/v1/auth/register",
        json={"username": username, "password": PASSWORD, "display_name": username},
        headers=headers(client),
    )
    assert response.status_code == 201, response.text
    return login(client, username)


class OfficeTestClient(FakeClient):
    """Deterministic tool-choice driver; not a model capability claim."""

    engine = None

    async def office_chat(self, snapshot, history, files, token, on_session):
        on_session("ses_office_test")
        source = next(f["id"] for f in files if f["name"] == "工作记录.xlsx")
        await self.engine.call(token, "read_table", {"file_id": source})
        validation = await self.engine.call(
            token,
            "validate_table",
            {
                "file_id": source,
                "rules": [
                    {"column": "人员", "required": True},
                    {"column": "工时", "type": "number", "minimum": 0.0},
                ],
            },
        )
        assert validation["valid"]
        summary = await self.engine.call(
            token,
            "transform_table",
            {
                "file_id": source,
                "operation": "group",
                "columns": ["人员"],
                "aggregations": {"工时": "sum"},
            },
        )
        report = await self.engine.call(
            token,
            "generate_document",
            {
                "title": "部门周报",
                "paragraphs": ["工作统计已核对。"],
                "table_file_id": summary["files"][0]["id"],
                "filename": "部门周报.docx",
            },
        )
        return "已根据工作记录完成汇总并生成周报。" + report["files"][0]["name"]


@pytest.fixture
def office(tmp_path):
    model = OfficeTestClient()
    app = create_app(
        Settings(state_dir=tmp_path, catalog_root=Path(__file__).resolve().parents[1]),
        client_factory=lambda _: model,
        source_factory=ExampleObservationSource,
    )
    with TestClient(app) as client:
        app.state.platform.initialize_admin("admin", "管理员", PASSWORD)
        model.engine = app.state.office
        login(client)
        yield app, client


def wait_execution(client, execution_id):
    for _ in range(200):
        response = client.get("/api/v1/executions/" + execution_id)
        assert response.status_code == 200, response.text
        value = response.json()
        if value["status"] not in ("queued", "running"):
            return value
        time.sleep(0.05)
    pytest.fail("office execution did not finish")


def upload(client, name="记录.csv", raw=b"name,hours\nA,2\nA,3\nB,4\n"):
    response = client.post("/api/v1/files", files={"file": (name, raw)}, headers=headers(client))
    assert response.status_code == 201, response.text
    return response.json()


def test_auth_registration_csrf_reset_and_isolation(office):
    app, alice = office
    bob = TestClient(app)
    assert bob.get("/api/v1/agents").status_code == 401
    bob_user = register(bob, "bob")
    assert bob.get("/api/v1/agents").status_code == 200
    assert (
        bob.patch(
            "/api/v1/tools/calculate", json={"enabled": False}, headers=headers(bob)
        ).status_code
        == 403
    )
    assert bob.post("/api/v1/auth/logout").status_code == 403
    changed = alice.patch(
        "/api/v1/admin/users/" + bob_user["id"],
        json={"password": "new-password-12345"},
        headers=headers(alice),
    )
    assert changed.status_code == 200
    assert bob.get("/api/v1/agents").status_code == 401
    login(bob, "bob", "new-password-12345")
    assert bob.get("/api/v1/agents").status_code == 403
    assert (
        bob.post(
            "/api/v1/auth/password",
            json={"current_password": "new-password-12345", "new_password": PASSWORD},
            headers=headers(bob),
        ).status_code
        == 200
    )
    login(bob, "bob")
    assert bob.get("/api/v1/agents").status_code == 200
    assert (
        alice.patch(
            "/api/v1/admin/users/" + bob_user["id"], json={"disabled": True}, headers=headers(alice)
        ).status_code
        == 200
    )
    assert bob.get("/api/v1/agents").status_code == 401
    with app.state.db.engine.connect() as c:
        assert c.execute(select(users.c.password_hash)).scalars().first().startswith("scrypt$")


def test_private_files_and_direct_tools(office):
    app, alice = office
    bob = TestClient(app)
    register(bob, "bob")
    file = upload(alice)
    assert bob.get(file["download_url"]).status_code == 404
    response = bob.post(
        "/api/v1/tool-executions",
        json={"tool": "read_table", "arguments": {"file_id": file["id"]}},
        headers=headers(bob),
    )
    assert response.status_code == 404
    body = {
        "tool": "transform_table",
        "arguments": {
            "file_id": file["id"],
            "operation": "group",
            "columns": ["name"],
            "aggregations": {"hours": "sum"},
        },
    }
    response = alice.post(
        "/api/v1/tool-executions", json=body, headers=headers(alice, "group-once")
    )
    assert response.status_code == 202, response.text
    execution = wait_execution(alice, response.json()["id"])
    assert execution["status"] == "succeeded", execution
    result = execution["files"][0]
    data = alice.get(result["download_url"]).content
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    assert list(workbook.active.values) == [("name", "hours_sum"), ("A", 5), ("B", 4)]
    workbook.close()
    replay = alice.post("/api/v1/tool-executions", json=body, headers=headers(alice, "group-once"))
    assert replay.json()["id"] == execution["id"]
    assert bob.get("/api/v1/executions/" + execution["id"]).status_code == 404
    assert bob.get(result["download_url"]).status_code == 404


def test_skill_releases_and_public_use(office):
    app, alice = office
    bob = TestClient(app)
    register(bob, "bob")
    body = {
        "title": "本部门流程",
        "description": "流程说明",
        "instructions": "读取文件并整理",
        "tools": ["read_document"],
    }
    skill = alice.post("/api/v1/skills", json=body, headers=headers(alice)).json()
    sid = skill["id"]
    assert bob.get("/api/v1/skills/" + sid).status_code == 404
    assert alice.post(f"/api/v1/skills/{sid}/publish", headers=headers(alice)).status_code == 200
    assert (
        alice.patch(
            f"/api/v1/skills/{sid}", json={"visibility": "public"}, headers=headers(alice)
        ).status_code
        == 200
    )
    assert bob.get(f"/api/v1/skills/{sid}").json()["editable"] is False
    assert bob.put(f"/api/v1/skills/{sid}", json=body, headers=headers(bob)).status_code == 403
    assert (
        alice.put(
            f"/api/v1/skills/{sid}",
            json={**body, "title": "未发布的新标题"},
            headers=headers(alice),
        ).status_code
        == 200
    )
    assert bob.get(f"/api/v1/skills/{sid}").json()["content"]["title"] == "本部门流程"
    agent = bob.post(
        "/api/v1/agents",
        json={
            "title": "我的助手",
            "description": "用于办公",
            "skills": [{"id": sid, "version": 1}],
        },
        headers=headers(bob),
    ).json()
    assert agent["content"]["tools"] == ["read_document"]
    aid = agent["id"]
    assert bob.post(f"/api/v1/agents/{aid}/publish", headers=headers(bob)).status_code == 200
    assert (
        alice.patch(
            f"/api/v1/skills/{sid}", json={"visibility": "private"}, headers=headers(alice)
        ).status_code
        == 200
    )
    assert (
        bob.post("/api/v1/conversations", json={"agent_id": aid}, headers=headers(bob)).status_code
        == 409
    )


def test_public_agent_cannot_publish_private_dependencies(office):
    _, client = office
    skill = client.post(
        "/api/v1/skills",
        json={"title": "私有技能", "description": "说明", "instructions": "处理"},
        headers=headers(client),
    ).json()
    client.post("/api/v1/skills/" + skill["id"] + "/publish", headers=headers(client))
    agent = client.post(
        "/api/v1/agents",
        json={
            "title": "助手",
            "description": "说明",
            "skills": [{"id": skill["id"], "version": 1}],
        },
        headers=headers(client),
    ).json()
    client.post("/api/v1/agents/" + agent["id"] + "/publish", headers=headers(client))
    assert (
        client.patch(
            "/api/v1/agents/" + agent["id"], json={"visibility": "public"}, headers=headers(client)
        ).status_code
        == 409
    )


def test_real_xlsx_to_docx_and_followup(office):
    app, client = office
    workbook = openpyxl.Workbook()
    workbook.active.append(["人员", "工时"])
    for row in [("张三", 2), ("张三", 3), ("李四", 4)]:
        workbook.active.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    file = upload(client, "工作记录.xlsx", stream.getvalue())
    conversation = client.post(
        "/api/v1/conversations",
        json={"agent_id": "builtin-agent-weekly-report"},
        headers=headers(client),
    ).json()
    cid = conversation["id"]
    response = client.post(
        f"/api/v1/conversations/{cid}/messages",
        json={"content": "请生成周报", "attachment_ids": [file["id"]]},
        headers=headers(client),
    )
    assert response.status_code == 202, response.text
    execution = wait_execution(client, response.json()["id"])
    assert execution["status"] == "succeeded", execution
    assert any("已加载 Skill" in e["label"] for e in execution["events"])
    docx = next(f for f in execution["files"] if f["name"].endswith(".docx"))
    document = Document(io.BytesIO(client.get(docx["download_url"]).content))
    assert "部门周报" in [p.text for p in document.paragraphs]
    row = [cell.text for cell in document.tables[0].rows[1].cells]
    assert row[0] == "张三" and float(row[1]) == 5.0
    response = client.post(
        f"/api/v1/conversations/{cid}/messages",
        json={"content": "请再生成一版"},
        headers=headers(client),
    )
    second = wait_execution(client, response.json()["id"])
    assert second["status"] == "succeeded", second
    assert client.get(docx["download_url"]).status_code == 200
    assert len(client.get(f"/api/v1/conversations/{cid}/messages").json()["items"]) == 4
    bob = TestClient(app)
    register(bob, "bob")
    assert bob.get(f"/api/v1/conversations/{cid}/messages").status_code == 404


def test_invalid_files_parameters_and_numeric_sandbox(office):
    _, client = office
    assert (
        client.post(
            "/api/v1/files", files={"file": ("payload.py", b"print('x')")}, headers=headers(client)
        ).status_code
        == 422
    )
    response = client.post(
        "/api/v1/tool-executions",
        json={"tool": "calculate", "arguments": {"expression": "__import__('os').listdir('/')"}},
        headers=headers(client),
    )
    result = wait_execution(client, response.json()["id"])
    assert result["status"] == "failed"
    response = client.post(
        "/api/v1/tool-executions",
        json={"tool": "calculate", "arguments": {"expression": "mean([2,4,6]) + 1"}},
        headers=headers(client),
    )
    assert wait_execution(client, response.json()["id"])["result"]["value"] == 5


def test_unverified_artifact_claim_cannot_complete_a_task(office):
    app, client = office

    async def hallucination(snapshot, history, files, token, on_session):
        return "已使用原工作记录重新生成新的周报文件，下载：https://example.invalid/fake.docx"

    app.state.client.office_chat = hallucination
    conversation = client.post(
        "/api/v1/conversations",
        json={"agent_id": "builtin-agent-weekly-report"},
        headers=headers(client),
    ).json()
    response = client.post(
        f"/api/v1/conversations/{conversation['id']}/messages",
        json={"content": "请再生成一份周报"},
        headers=headers(client),
    )
    result = wait_execution(client, response.json()["id"])
    assert result["status"] == "failed"
    assert not result["files"]
    assert "未登记真实产物" in result["error"]
    messages = client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()["items"]
    assert [m["role"] for m in messages] == ["user"]


def test_registered_files_replace_model_download_links():
    from types import SimpleNamespace

    from apps.orchestrator.office import verified_reply

    reply = verified_reply(
        "文档已生成，请下载 https://example.invalid/invented.docx?signature=fake",
        [SimpleNamespace(name="实际周报.docx")],
    )
    assert "实际周报.docx" in reply and "https://" not in reply
    assert verified_reply("请先上传工作记录，再生成周报。", []) == "请先上传工作记录，再生成周报。"
    assert verified_reply("已读取文件，下面是分析结论。", []) == "已读取文件，下面是分析结论。"


def test_document_content_preserves_body_and_markdown_blocks(office):
    _, client = office
    response = client.post(
        "/api/v1/tool-executions",
        json={"tool": "generate_document", "arguments": {
            "title": "内容总结", "filename": "总结.md",
            "content": "# 主题\n\n这是第一段实际正文。\n\n## 要点\n\n- 要点一\n- 要点二\n\n1. 执行一步\n\n| 人员 | 工时 |\n|---|---|\n| 张三 | 5 |\n\n这是最后一段。",
        }},
        headers=headers(client),
    )
    assert response.status_code == 202, response.text
    execution = wait_execution(client, response.json()["id"])
    assert execution["status"] == "succeeded", execution
    file = execution["files"][0]
    assert file["name"] == "总结.docx"
    doc = Document(io.BytesIO(client.get(file["download_url"]).content))
    assert "这是第一段实际正文。" in [p.text for p in doc.paragraphs]
    assert "这是最后一段。" in [p.text for p in doc.paragraphs]
    assert any(p.style.name == "List Bullet" and p.text == "要点一" for p in doc.paragraphs)
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == [["人员", "工时"], ["张三", "5"]]
    assert execution["result"]["body_characters"] > 20


@pytest.mark.parametrize("arguments", [
    {"title": "只有标题"},
    {"title": "空格正文", "content": " \n\t ", "paragraphs": [" "]},
    {"title": "空段落", "sections": [{"kind": "paragraph", "text": " "}]},
    {"title": "空表格", "sections": [{"kind": "table", "rows": [[None, " "]]}]},
])
def test_title_only_document_is_rejected_before_execution(office, arguments):
    app, client = office
    response = client.post(
        "/api/v1/tool-executions",
        json={"tool": "generate_document", "arguments": arguments},
        headers=headers(client),
    )
    assert response.status_code == 422
    assert "Word 正文不能为空" in response.json()["error"]["message"]
    from packages.adapters.platform_store import files

    with app.state.db.engine.connect() as c:
        assert not c.execute(select(files)).first()


@pytest.mark.parametrize("body", ["", "这是模板中的实际正文。"])
def test_template_document_requires_actual_body(office, body):
    _, client = office
    doc = Document()
    doc.add_heading("模板标题", 0)
    if body:
        doc.add_paragraph(body)
    buffer = io.BytesIO()
    doc.save(buffer)
    file = upload(client, "模板.docx", buffer.getvalue())
    response = client.post(
        "/api/v1/tool-executions",
        json={"tool": "generate_document", "arguments": {"title": "总结", "template_file_id": file["id"]}},
        headers=headers(client),
    )
    result = wait_execution(client, response.json()["id"])
    if body:
        assert result["status"] == "succeeded", result
        repaired = Document(io.BytesIO(client.get(result["files"][0]["download_url"]).content))
        assert body in [p.text for p in repaired.paragraphs]
    else:
        assert result["status"] == "failed" and not result["files"]
        assert "Word 正文为空" in result["error"]


def test_scope_cannot_access_other_task_file(office):
    app, client = office
    file = upload(client)
    token = "invalid-capability"
    response = client.post(
        "/internal/v1/office/read_table",
        json={"context_token": token, "arguments": {"file_id": file["id"]}},
    )
    assert response.status_code == 403

    # Even possession of the service credential is insufficient without a task capability.
    app.state.platform_service_token = "service-test-credential"
    response = client.post(
        "/internal/v1/office/read_table",
        json={"context_token": token, "arguments": {"file_id": file["id"]}},
        headers={"X-MCP-Service-Token": "service-test-credential"},
    )
    assert response.status_code == 403



def test_upload_replay_and_conflict(office):
    _, client = office
    request_headers = headers(client, "same-upload-request")
    first = client.post(
        "/api/v1/files", files={"file": ("记录.csv", b"name,value\nA,1")}, headers=request_headers
    )
    again = client.post(
        "/api/v1/files", files={"file": ("记录.csv", b"name,value\nA,1")}, headers=request_headers
    )
    assert first.status_code == again.status_code == 201
    assert first.json()["id"] == again.json()["id"]
    conflict = client.post(
        "/api/v1/files", files={"file": ("记录.csv", b"name,value\nA,2")}, headers=request_headers
    )
    assert conflict.status_code == 409


def test_office_cancel_and_queue_survives(tmp_path):
    import asyncio

    app = create_app(
        Settings(state_dir=tmp_path),
        client_factory=FakeClient,
        source_factory=ExampleObservationSource,
    )
    with TestClient(app) as client:
        app.state.platform.initialize_admin("admin", "管理员", PASSWORD)
        login(client)
        original = app.state.office.runner.run
        stopped = []

        async def delayed(*args, **kwargs):
            try:
                await asyncio.sleep(30)
            finally:
                stopped.append(True)

        app.state.office.runner.run = delayed
        response = client.post(
            "/api/v1/tool-executions",
            json={"tool": "calculate", "arguments": {"expression": "1+1"}},
            headers=headers(client),
        )
        execution_id = response.json()["id"]
        for _ in range(50):
            if client.get("/api/v1/executions/" + execution_id).json()["status"] == "running":
                break
            time.sleep(0.02)
        response = client.post(f"/api/v1/executions/{execution_id}/cancel", headers=headers(client))
        assert response.status_code == 200 and response.json()["status"] == "cancelled"
        for _ in range(50):
            if stopped:
                break
            time.sleep(0.02)
        assert stopped
        app.state.office.runner.run = original
        response = client.post(
            "/api/v1/tool-executions",
            json={"tool": "calculate", "arguments": {"expression": "6*7"}},
            headers=headers(client),
        )
        result = wait_execution(client, response.json()["id"])
        assert result["status"] == "succeeded" and result["result"]["value"] == 42


def test_office_restart_and_valid_scope_file_boundary(tmp_path):
    import hashlib

    from apps.orchestrator.office import OfficeEngine
    from packages.domain.core import DomainError

    app = create_app(Settings(state_dir=tmp_path), client_factory=FakeClient, worker_enabled=False)
    with TestClient(app) as client:
        account = app.state.platform.initialize_admin("admin", "管理员", PASSWORD)
        login(client)
        first, second = upload(client), upload(client, "另一份.csv")
        with app.state.db.engine.begin() as c:
            execution_id = app.state.office.create(
                account["id"],
                "tool",
                {"tool": "calculate", "arguments": {"expression": "1"}},
                None,
                c,
            )
        app.state.office.update(execution_id, "running")
        token = "test-valid-scoped-capability"
        app.state.office.scopes[hashlib.sha256(token.encode()).hexdigest()] = dict(
            execution_id=execution_id,
            agent_id="builtin-agent-general-assistant",
            version=1,
            user=account["id"],
            conversation_id=None,
            files={first["id"]},
            tools={"read_table"},
            expires=time.monotonic() + 60,
            revoked=False,
            calls=0,
            results={},
        )
        with pytest.raises(DomainError) as denied:
            client.portal.call(
                app.state.office.call, token, "read_table", {"file_id": second["id"]}
            )
        assert denied.value.code == "FILE_FORBIDDEN"
        restarted = OfficeEngine(app.state.platform, app.state.client)
        assert restarted.info(execution_id, account["id"]).status == "interrupted"
        assert not restarted.scopes


def test_general_assistant_upgrades_existing_resources_and_keeps_private_history(office):
    from packages.adapters.office_resources import ResourceService
    from packages.adapters.platform_store import resources, versions

    app, client = office
    admin = client.get("/api/v1/auth/me").json()["user"]["id"]
    general_id = "builtin-agent-general-assistant"
    old_agents = {
        item["id"]: item
        for item in client.get("/api/v1/agents").json()["items"]
        if item["id"] != general_id
    }
    legacy = client.post(
        "/api/v1/chat/sessions", json={"schema_version": "1.0"}, headers=headers(client)
    ).json()
    old_reply = client.post(
        f"/api/v1/chat/sessions/{legacy['session_id']}/messages",
        json={"schema_version": "1.0", "message": "这是升级前保存的普通对话"},
        headers=headers(client),
    )
    assert old_reply.status_code == 200
    old_messages = client.get(f"/api/v1/chat/sessions/{legacy['session_id']}/messages").json()
    # A deployed office database can already contain all three original assistants.
    ids = [general_id, "builtin-skill-general-assistant"]
    with app.state.db.engine.begin() as c:
        c.execute(versions.delete().where(versions.c.resource_id.in_(ids)))
        c.execute(resources.delete().where(resources.c.id.in_(ids)))
    disabled = client.patch(
        "/api/v1/tools/calculate", json={"enabled": False}, headers=headers(client)
    )
    assert disabled.status_code == 200
    service = ResourceService(app.state.platform)
    service.seed(admin)
    first = service.get(general_id, admin)
    service.seed(admin)
    assert service.get(general_id, admin) == first
    assert first.builtin and first.visibility == "public" and not first.editable
    assert first.content.title == "通用助手" and first.published_version == 1
    for item in client.get("/api/v1/agents").json()["items"]:
        if item["id"] in old_agents:
            assert item == old_agents[item["id"]]

    requests = []

    async def plain_reply(snapshot, history, files, token, on_session):
        assert snapshot["id"] == general_id and snapshot["skills"]
        assert not files
        assert all(not name.startswith("industrial") for name in snapshot["content"]["tools"])
        assert "calculate" not in snapshot["content"]["tools"]
        requests.append([m["content"] for m in history])
        on_session("ses_general_test")
        return "可以先明确目的、参与人员和时间安排。"

    app.state.client.office_chat = plain_reply
    conversation = client.post(
        "/api/v1/conversations", json={"agent_id": general_id}, headers=headers(client)
    ).json()
    for content in ["帮我整理会议议程", "请检查 PID 状态"]:
        response = client.post(
            f"/api/v1/conversations/{conversation['id']}/messages",
            json={"content": content}, headers=headers(client),
        )
        assert response.status_code == 202
        result = wait_execution(client, response.json()["id"])
        assert result["status"] == "succeeded" and not result["files"]
    assert len(requests) == 2 and requests[0][0] in requests[1]
    assert not app.state.client.calls or all(call[0] == "reception" for call in app.state.client.calls)
    assert not client.get("/api/v1/runs").json()["items"]
    assert not client.get("/api/v1/tuning-runs").json()["items"]
    assert client.get(f"/api/v1/chat/sessions/{legacy['session_id']}/messages").json() == old_messages
    assert len(client.get(f"/api/v1/conversations/{conversation['id']}/messages").json()["items"]) == 4
    member = TestClient(app)
    register(member, "general-member")
    assert general_id in {r["id"] for r in member.get("/api/v1/agents").json()["items"]}
    assert member.get(f"/api/v1/conversations/{conversation['id']}/messages").status_code == 404
    assert member.get(f"/api/v1/chat/sessions/{legacy['session_id']}/messages").status_code == 404
    assert not member.get("/api/v1/chat/sessions").json()["items"]
