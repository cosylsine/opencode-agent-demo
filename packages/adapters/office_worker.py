"""One bounded subprocess per office operation. No shell, network or user code."""

import ast
import csv
import datetime as dt
import io
import json
import math
import operator
import os
import re
import resource
import statistics
import sys
import zipfile
from pathlib import Path


def document_content(doc, content, add_table):
    """Write text and supported Markdown blocks; never treat it as a file path."""
    lines, paragraph = content.splitlines(), []

    def flush():
        if paragraph:
            doc.add_paragraph("\n".join(paragraph))
            paragraph.clear()

    def cells(line):
        return [v.strip().replace(r"\|", "|") for v in re.split(r"(?<!\\)\|", line.strip().strip("|"))]

    i, code = 0, False
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith(("```", "~~~")):
            flush()
            code = not code
        elif code:
            doc.add_paragraph(lines[i])
        elif not line:
            flush()
        elif "|" in line and i + 1 < len(lines) and all(
            re.fullmatch(r":?-{3,}:?", c) for c in cells(lines[i + 1])
        ):
            flush()
            rows = [cells(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(cells(lines[i]))
                i += 1
            add_table(rows)
            continue
        elif match := re.match(r"^(#{1,6})\s+(.+)$", line):
            flush()
            doc.add_heading(match[2], min(len(match[1]), 3))
        elif match := re.match(r"^[-*+]\s+(.+)$", line):
            flush()
            doc.add_paragraph(match[1], style="List Bullet")
        elif match := re.match(r"^\d+[.)]\s+(.+)$", line):
            flush()
            doc.add_paragraph(match[1], style="List Number")
        else:
            paragraph.append(lines[i])
        i += 1
    flush()


def checked_archive(path):
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as z:
            if (
                len(z.infolist()) > 20000
                or sum(i.file_size for i in z.infolist()) > 200 * 1024 * 1024
            ):
                raise ValueError("文件解压内容过大")


def scalar(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)


def table(path, sheet=None):
    path = Path(path)
    warnings, sheets = [], []
    if path.suffix == ".xlsx":
        import openpyxl

        checked_archive(path)
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        sheets = workbook.sheetnames
        if sheet and sheet not in sheets:
            workbook.close()
            raise ValueError("工作表不存在，可选：" + "、".join(sheets))
        ws = workbook[sheet] if sheet else workbook.active
        if ws.max_row and ws.max_row > 100001 or ws.max_column and ws.max_column > 200:
            workbook.close()
            raise ValueError("表格超过 100000 行或 200 列，请先拆分")
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([scalar(v) for v in row])
            if len(rows) > 100001:
                raise ValueError("表格超过 100000 行")
        workbook.close()
        warnings = ["公式按已有缓存值读取，未重新计算；空单元格可能是公式缺少缓存。"]
    elif path.suffix == ".csv":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        rows = list(csv.reader(io.StringIO(text)))
        if len(rows) > 100001 or any(len(r) > 200 for r in rows):
            raise ValueError("表格超过 100000 行或 200 列")
    else:
        raise ValueError("此工具需要 XLSX 或 CSV")
    if not rows:
        raise ValueError("表格为空")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("表头必须非空且不能重复，请先整理首行字段")
    values = [
        dict(zip(headers, r + [None] * max(0, len(headers) - len(r))))
        for r in rows[1:]
        if any(v not in (None, "") for v in r)
    ]
    return headers, values, sheets, warnings


def number(value):
    if isinstance(value, bool):
        raise ValueError("布尔值不能作为数值")
    value = float(value)
    if not math.isfinite(value):
        raise ValueError("数值必须有限")
    return value


def expression(source):
    tree = ast.parse(source, mode="eval")
    if len(list(ast.walk(tree))) > 250:
        raise ValueError("表达式过长")
    binary = {
        ast.Add: operator.add,
        ast.Sub: operator.sub,
        ast.Mult: operator.mul,
        ast.Div: operator.truediv,
        ast.Mod: operator.mod,
        ast.Pow: operator.pow,
    }
    functions = {
        "sum": sum,
        "mean": statistics.mean,
        "min": min,
        "max": max,
        "round": round,
        "abs": abs,
    }

    def evaluate(node):
        if isinstance(node, ast.Expression):
            return evaluate(node.body)
        if isinstance(node, ast.Constant) and type(node.value) in (int, float):
            return number(node.value)
        if isinstance(node, (ast.List, ast.Tuple)):
            return [evaluate(n) for n in node.elts]
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            return (-1 if isinstance(node.op, ast.USub) else 1) * number(evaluate(node.operand))
        if isinstance(node, ast.BinOp) and type(node.op) in binary:
            left, right = number(evaluate(node.left)), number(evaluate(node.right))
            if isinstance(node.op, ast.Pow) and (abs(right) > 100 or abs(left) > 1e20):
                raise ValueError("指数超出计算范围")
            return number(binary[type(node.op)](left, right))
        if (
            isinstance(node, ast.Call)
            and isinstance(node.func, ast.Name)
            and node.func.id in functions
            and not node.keywords
        ):
            args = [evaluate(n) for n in node.args]
            if node.func.id == "round" and len(args) == 2:
                args[1] = int(args[1])
            return number(functions[node.func.id](*args))
        raise ValueError("表达式包含不支持的操作")

    return evaluate(tree)


def safe_cell(value):
    # All generated strings stay literal, including text beginning with '='.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return scalar(value)


def spreadsheet(output, filename, headers, rows):
    suffix = ".csv" if filename.lower().endswith(".csv") else ".xlsx"
    path = output / ("result" + suffix)
    if suffix == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows([[safe_cell(v) for v in r] for r in rows])
    else:
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = "结果"
        ws.append([safe_cell(h) for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B66")
        for row in rows:
            ws.append([safe_cell(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        workbook.save(path)
        workbook.close()
    name = Path(filename.replace("\\", "/")).stem[:100] + suffix
    return {"path": str(path), "name": name}


def document_text(path):
    path = Path(path)
    if path.suffix == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph

        checked_archive(path)
        doc = Document(path)
        text = "\n".join(
            item.text
            if isinstance(item, Paragraph)
            else "\n".join("\t".join(c.text for c in row.cells) for row in item.rows)
            for item in doc.iter_inner_content()
        )
    elif path.suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("暂不支持加密 PDF")
        if len(reader.pages) > 1000:
            raise ValueError("PDF 超过 1000 页，请先拆分")
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
        if not text.strip():
            raise ValueError("未提取到文本；扫描 PDF 暂不支持，请提供文本型文件")
    elif path.suffix in (".txt", ".md"):
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
    else:
        raise ValueError("不支持此文档格式")
    return text


def execute(tool, a, paths, output):
    from packages.adapters.office_extra_worker import EXTRA_TOOLS, execute_extra

    if tool in EXTRA_TOOLS:
        return execute_extra(tool, a, paths, output)
    if tool == "calculate":
        return {"value": expression(a["expression"])}, []
    if tool == "date_calculate":
        date = dt.date.fromisoformat(a["date"])
        return (
            {"days": (dt.date.fromisoformat(a["other_date"]) - date).days}
            if a["operation"] == "difference"
            else {"date": (date + dt.timedelta(days=a["days"])).isoformat()}
        ), []
    if tool == "read_document":
        text = document_text(paths["file_id"])
        end = a["start"] + a["limit"]
        return {
            "text": text[a["start"] : end],
            "total_characters": len(text),
            "next_start": end if end < len(text) else None,
        }, []
    if tool in ("read_table", "transform_table", "validate_table"):
        headers, rows, sheets, warnings = table(paths["file_id"], a.get("sheet"))
        if tool == "read_table":
            if a.get("columns"):
                if any(c not in headers for c in a["columns"]):
                    raise ValueError("选择的字段不存在")
                headers = a["columns"]
                rows = [{c: r[c] for c in headers} for r in rows]
            end = a["offset"] + a["limit"]
            return dict(
                headers=headers,
                rows=rows[a["offset"] : end],
                total_rows=len(rows),
                sheets=sheets,
                warnings=warnings,
                next_offset=end if end < len(rows) else None,
            ), []
        if tool == "validate_table":
            issues = []
            total = 0
            for rule in a["rules"]:
                column, seen = rule["column"], set()
                if column not in headers:
                    issues.append({"row": 1, "column": column, "issue": "缺少字段"})
                    total += 1
                    continue
                for i, row in enumerate(rows, 2):
                    value, errors = row[column], []
                    if value in (None, ""):
                        if rule["required"]:
                            errors.append("必填值缺失")
                    else:
                        if rule["type"] == "text" and not isinstance(value, str):
                            errors.append("类型不符合要求")
                        try:
                            if (
                                rule["type"] == "number"
                                or rule["minimum"] is not None
                                or rule["maximum"] is not None
                            ):
                                n = number(value)
                                if (
                                    rule["minimum"] is not None
                                    and n < rule["minimum"]
                                    or rule["maximum"] is not None
                                    and n > rule["maximum"]
                                ):
                                    errors.append("数值越界")
                            if rule["type"] == "date":
                                dt.date.fromisoformat(str(value)[:10])
                        except (ValueError, TypeError):
                            errors.append("类型不符合要求")
                        if rule["unique"] and str(value) in seen:
                            errors.append("值重复")
                        seen.add(str(value))
                    for err in errors:
                        total += 1
                        if len(issues) < 100:
                            issues.append({"row": i, "column": column, "issue": err})
            return dict(
                valid=total == 0,
                issue_count=total,
                issues=issues,
                total_rows=len(rows),
                warnings=warnings,
            ), []
        columns, operation = a["columns"], a["operation"]
        if any(c not in headers for c in columns + list(a["aggregations"])):
            raise ValueError("指定字段不存在，可选：" + "、".join(headers))
        if operation == "filter":

            def matches(row, f):
                if f["column"] not in headers:
                    raise ValueError("筛选字段不存在")
                value, expected, op = row[f["column"]], f["value"], f["operator"]
                if op == "empty":
                    return value in (None, "")
                if op == "contains":
                    return str(expected) in str(value or "")
                if op in ("eq", "ne"):
                    equal = str(value) == str(expected)
                    if type(value) in (int, float) or type(expected) in (int, float):
                        try:
                            equal = number(value) == number(expected)
                        except (ValueError, TypeError):
                            equal = False
                    return equal if op == "eq" else not equal
                return {"gt": operator.gt, "ge": operator.ge, "lt": operator.lt, "le": operator.le}[
                    op
                ](number(value), number(expected))

            rows = [r for r in rows if all(matches(r, f) for f in a["filters"])]
        elif operation == "sort":
            if not columns:
                raise ValueError("请选择排序列")

            def sort_key(r):
                result = []
                for c in columns:
                    v = r[c]
                    try:
                        result.append((0, number(v)))
                    except (TypeError, ValueError):
                        result.append((1, str(v or "")))
                return result

            rows.sort(key=sort_key, reverse=a["descending"])
        elif operation == "deduplicate":
            seen, unique = set(), []
            for row in rows:
                key = tuple(str(row[c]) for c in columns or headers)
                if key not in seen:
                    seen.add(key)
                    unique.append(row)
            rows = unique
        elif operation == "group":
            if not a["aggregations"]:
                raise ValueError("请指定统计列及方法")
            grouped = {}
            for row in rows:
                grouped.setdefault(tuple(row[c] for c in columns), []).append(row)
            result = []
            for key, items in grouped.items():
                out = dict(zip(columns, key))
                for column, op in a["aggregations"].items():
                    values = (
                        [number(r[column]) for r in items if r[column] not in (None, "")]
                        if op != "count"
                        else []
                    )
                    out[column + "_" + op] = (
                        len(items)
                        if op == "count"
                        else (
                            {"sum": sum, "mean": statistics.mean, "min": min, "max": max}[op](
                                values
                            )
                            if values
                            else None
                        )
                    )
                result.append(out)
            headers, rows = columns + [c + "_" + op for c, op in a["aggregations"].items()], result
        elif operation in ("join", "compare"):
            if not a["other_file_id"] or not columns:
                raise ValueError("请选择另一张表及匹配列")
            right_headers, right, _, _ = table(paths["other_file_id"])
            if any(c not in right_headers for c in columns):
                raise ValueError("另一张表缺少匹配字段")

            def key(r):
                return tuple(str(r[c]) for c in columns)

            indexed = {}
            for row in right:
                k = key(row)
                if k in indexed:
                    raise ValueError("另一张表匹配键重复，请先去重")
                indexed[k] = row
            if operation == "join":
                added = {
                    c: c + "_另一表" if c in headers else c
                    for c in right_headers
                    if c not in columns
                }
                rows = [
                    {
                        **r,
                        **{
                            target: indexed.get(key(r), {}).get(source)
                            for source, target in added.items()
                        },
                    }
                    for r in rows
                ]
                headers += list(added.values())
            else:
                result = []
                left_keys = set()
                for row in rows:
                    k = key(row)
                    left_keys.add(k)
                    other = indexed.get(k)
                    changed = [
                        c
                        for c in headers
                        if c not in columns
                        and other is not None
                        and str(row[c]) != str(other.get(c))
                    ]
                    if other is None or changed:
                        result.append(
                            {
                                **row,
                                "差异": "仅本表存在"
                                if other is None
                                else "字段不同：" + "、".join(changed),
                            }
                        )
                for k, row in indexed.items():
                    if k not in left_keys:
                        result.append({**{c: row.get(c) for c in headers}, "差异": "仅另一表存在"})
                rows, headers = result, headers + ["差异"]
        artifact = spreadsheet(
            output, a["filename"], headers, [[r.get(c) for c in headers] for r in rows]
        )
        return dict(headers=headers, rows=rows[:30], total_rows=len(rows), warnings=warnings), [
            artifact
        ]
    if tool == "generate_spreadsheet":
        if not a["headers"] or len(a["headers"]) > 200 or len(a["rows"]) > 100000:
            raise ValueError("表格字段或行数超出范围")
        return {"rows": len(a["rows"])}, [
            spreadsheet(output, a["filename"], a["headers"], a["rows"])
        ]
    if tool == "export_text":
        suffix = ".txt" if a["filename"].lower().endswith(".txt") else ".md"
        path = output / ("result" + suffix)
        path.write_text(a["content"], encoding="utf-8")
        return {"characters": len(a["content"])}, [
            {"path": str(path), "name": Path(a["filename"].replace("\\", "/")).stem[:100] + suffix}
        ]
    if tool == "generate_document":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt

        if a["template_file_id"]:
            if Path(paths["template_file_id"]).suffix != ".docx":
                raise ValueError("模板须为 DOCX")
            checked_archive(paths["template_file_id"])
        doc = Document(paths["template_file_id"]) if a["template_file_id"] else Document()
        style = doc.styles["Normal"]
        style.font.name, style.font.size = "宋体", Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if a["title"]:
            doc.add_heading(a["title"], 0)

        def add_table(rows):
            if not rows:
                return
            t = doc.add_table(rows=0, cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for row in rows:
                cells = t.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = "" if value is None else str(value)

        document_content(doc, a["content"], add_table)
        for paragraph in a["paragraphs"]:
            doc.add_paragraph(paragraph)
        for section in a["sections"]:
            if section["kind"] == "heading":
                doc.add_heading(section["text"], section["level"])
            elif section["kind"] == "paragraph":
                doc.add_paragraph(section["text"])
            elif section["kind"] == "list":
                for item in section["items"]:
                    doc.add_paragraph(item, style="List Bullet")
            else:
                add_table(section["rows"])
        if a["table_file_id"]:
            headers, rows, _, _ = table(paths["table_file_id"])
            if len(rows) > 5000:
                raise ValueError("附表超过 5000 行，请先汇总")
            add_table([headers] + [[r[c] for c in headers] for r in rows])
        body = [p.text for p in doc.paragraphs if p.style.name != "Title" and p.text.strip()]
        cell_text = [cell.text for t in doc.tables for row in t.rows for cell in row.cells if cell.text.strip()]
        body_characters = sum(len(text.strip()) for text in [*body, *cell_text])
        if not body_characters:
            raise ValueError("Word 正文为空，未生成文件；请提供实际段落、列表或表格内容，不能只传标题。")
        path = output / "result.docx"
        doc.save(path)
        return {"title": a["title"], "body_characters": body_characters,
                "paragraph_count": len(body), "table_count": len(doc.tables)}, [
            {"path": str(path), "name": Path(a["filename"].replace("\\", "/")).stem[:100] + ".docx"}
        ]
    raise ValueError("未知工具")


def main():
    resource.setrlimit(resource.RLIMIT_CPU, (45, 45))
    resource.setrlimit(resource.RLIMIT_AS, (1536 * 1024 * 1024, 1536 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_FSIZE, (64 * 1024 * 1024, 64 * 1024 * 1024))
    os.umask(0o077)
    request = json.load(sys.stdin)
    try:
        result, artifacts = execute(
            request["tool"], request["arguments"], request["paths"], Path(request["output"])
        )
        print(
            json.dumps(
                {"result": result, "artifacts": artifacts}, ensure_ascii=False, allow_nan=False
            )
        )
    except Exception as exc:
        print(
            json.dumps(
                {
                    "error": str(exc)[:300]
                    if isinstance(exc, (ValueError, TypeError, KeyError, ZeroDivisionError))
                    else "文件解析或生成失败，请检查格式和内容"
                },
                ensure_ascii=False,
            )
        )


if __name__ == "__main__":
    main()
