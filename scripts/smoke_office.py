"""Real office stack acceptance. TEST protocol runs must be explicitly labelled."""

import argparse
import asyncio
import hashlib
import io
import json
import os
import sys
import uuid
from pathlib import Path

import httpx
import openpyxl
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.acceptance_login import acceptance_login


async def main(base, output, protocol_only=False):
    evidence = {
        "schema_version": "1.0",
        "model_validation": "TEST_PROTOCOL_ONLY" if protocol_only else "REMOTE_MODEL",
        "real_qwen_inference_verified": not protocol_only,
        "turns": [],
    }
    async with httpx.AsyncClient(base_url=base, trust_env=False, timeout=30) as client:
        csrf = await acceptance_login(client, base)

        def headers():
            return {"Origin": base, "X-CSRF-Token": csrf, "Idempotency-Key": str(uuid.uuid4())}

        book = openpyxl.Workbook()
        book.active.append(["人员", "工时"])
        for row in [("张三", 2), ("张三", 3), ("李四", 4)]:
            book.active.append(row)
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        response = await client.post(
            "/api/v1/files", files={"file": ("工作记录.xlsx", buffer.getvalue())}, headers=headers()
        )
        response.raise_for_status()
        file_id = response.json()["id"]
        # Use a newly authored Skill and assistant so this also verifies publishing and binding.
        response = await client.post(
            "/api/v1/skills",
            json={
                "title": "办公闭环验收技能",
                "description": "核对人员工时并生成周报。",
                "inputs": "XLSX，首行为人员和工时。",
                "instructions": (
                    "每轮都读取原始工作记录，校验人员必填和工时非负，"
                    "使用 transform_table 按人员分组、对工时求和。"
                    "生成 Word 周报时，必须把本轮 transform_table 返回的文件 ID "
                    "传给 generate_document.table_file_id，直接插入真实 Word 汇总表；"
                    "不要手工重抄汇总值，也不要用段落或 Markdown 管道文本代替表格。"
                    "正文简要说明数据来源和校验结果。后续轮次保留原文件及既有产物，"
                    "重新生成汇总文件和周报。"
                ),
                "tools": ["read_table", "validate_table", "transform_table", "generate_document"],
            },
            headers=headers(),
        )
        response.raise_for_status()
        skill = response.json()
        response = await client.post(
            "/api/v1/skills/" + skill["id"] + "/publish", headers=headers()
        )
        response.raise_for_status()
        response = await client.post(
            "/api/v1/agents",
            json={
                "title": "办公闭环验收助手",
                "description": "按自定义技能生成周报。",
                "skills": [{"id": skill["id"], "version": 1}],
            },
            headers=headers(),
        )
        response.raise_for_status()
        agent = response.json()
        response = await client.post(
            "/api/v1/agents/" + agent["id"] + "/publish", headers=headers()
        )
        response.raise_for_status()
        response = await client.post(
            "/api/v1/conversations", json={"agent_id": agent["id"]}, headers=headers()
        )
        response.raise_for_status()
        conversation_id = response.json()["id"]
        for turn in range(2):
            body = {
                "content": "请按技能读取、校验工作记录，按人员汇总工时，生成带汇总表的 Word 周报。"
                if turn == 0
                else "请使用原工作记录再生成一份周报，保留原文件。",
                "attachment_ids": [file_id] if turn == 0 else [],
            }
            response = await client.post(
                f"/api/v1/conversations/{conversation_id}/messages", json=body, headers=headers()
            )
            response.raise_for_status()
            execution_id = response.json()["id"]
            for _ in range(310):
                response = await client.get("/api/v1/executions/" + execution_id)
                response.raise_for_status()
                run = response.json()
                if run["status"] not in ("queued", "running"):
                    break
                await asyncio.sleep(2)
            if run["status"] != "succeeded":
                raise RuntimeError("办公任务未完成：" + str(run.get("error")))
            tools = {e.get("tool") for e in run["events"] if e.get("tool")}
            required = {"read_table", "validate_table", "transform_table", "generate_document"}
            if not required.issubset(tools) or not any(
                "已加载 Skill" in e["label"] for e in run["events"]
            ):
                raise RuntimeError("缺少技能加载或必要工具的执行证据")
            file = next(f for f in run["files"] if f["name"].endswith(".docx"))
            document_response = await client.get(file["download_url"])
            document_response.raise_for_status()
            document = Document(io.BytesIO(document_response.content))
            rows = [[c.text for c in row.cells] for table in document.tables for row in table.rows]
            if not all(
                any(person in row and any(value in (str(total), f"{total}.0") for value in row) for row in rows)
                for person, total in (("张三", 5), ("李四", 4))
            ):
                raise RuntimeError("生成文档缺少已核对的汇总结果")
            evidence["turns"].append(
                {
                    "execution_id": execution_id,
                    "tools": sorted(tools),
                    "file_id": file["id"],
                    "sha256": hashlib.sha256(document_response.content).hexdigest(),
                    "docx_table_checked": True,
                }
            )
        previous = evidence["turns"][0]["file_id"]
        previous_response = await client.get(f"/api/v1/files/{previous}/download")
        previous_response.raise_for_status()
        assert hashlib.sha256(previous_response.content).hexdigest() == evidence["turns"][0]["sha256"]
        assert evidence["turns"][1]["file_id"] != previous
        original_response = await client.get(f"/api/v1/files/{file_id}/download")
        original_response.raise_for_status()
        assert original_response.content == buffer.getvalue()
        evidence.update(
            conversation_id=conversation_id,
            skill_id=skill["id"],
            agent_id=agent["id"],
            original_and_previous_files_unchanged=True,
            model=(await client.get("/api/v1/models")).json(),
            success=True,
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print("办公闭环验收通过：两轮真实文件已核对；报告：" + str(output))
    if protocol_only:
        print("本次使用 TEST 协议服务，不代表真实 Qwen 推理能力。")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", default="http://127.0.0.1:3000")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(os.getenv("DEMO_STATE_DIR", "var/backend")) / "office-acceptance.json",
    )
    parser.add_argument("--protocol-only", action="store_true")
    args = parser.parse_args()
    asyncio.run(main(args.base, args.output, args.protocol_only))
